"""Regression checks for a real fleet fuel sheet: timestamp dates, empty cost, several vehicles over several years."""
import io
import itertools
from datetime import datetime

import openpyxl
import pytest
from PIL import Image

from app.extraction.extractor import normalize_field
from app.extraction.normalizer import normalize_date
from app.extraction.pipeline import classify_document
from app.extraction.registry import resolve_schema
from app.reporting.analytics import analyze_records, parse_period
from app.reporting.charts import PALETTE, render_monthly_chart
from app.reporting.excel_exporter import create_styled_excel_report
from app.reporting.platform_excel_exporter import TREND_SHEET


def fuel_schema():
    return resolve_schema("fuel_expense", "fleet_fuel_expenses")


def fuel_records(n_plates=6, years=(2025, 2026), months=range(1, 10), timestamp=True, with_cost=False):
    """Records built through the real field normalizer, like an Excel sheet read row by row (no document-specific data)."""
    schema = fuel_schema()
    plates = [f"ZZ{100 + i}XX" for i in range(n_plates)]
    out = []
    for k, ((y, m), plate) in enumerate(itertools.product([(y, m) for y in years for m in months], plates)):
        when = datetime(y, m, 1)
        out.append({"source_document": "foglio.xlsx", "source_page": 1, "fields": {
            "targa": normalize_field(schema.fields["targa"], plate),
            "data_rifornimento": normalize_field(schema.fields["data_rifornimento"],
                                                 str(when) if timestamp else when.date().isoformat()),
            "litri": normalize_field(schema.fields["litri"], 100 + (k % 7) * 20.5),
            "importo_eur": normalize_field(schema.fields["importo_eur"], 50.0 + k if with_cost else None),
            "chilometri": normalize_field(schema.fields["chilometri"], 1000 + k),
        }})
    return out


@pytest.mark.offline
@pytest.mark.parametrize("raw,iso", [
    ("2025-01-01 00:00:00", "2025-01-01"), ("2025-01-01 00:00:00.000", "2025-01-01"),
    ("2025-01-01T08:30:00Z", "2025-01-01"), ("2025-01-01 00:00:00+01:00", "2025-01-01"),
    ("04/11/2026 14:30", "2026-11-04"), ("31/07/2026", "2026-07-31"), ("15 October 2026", "2026-10-15"),
])
def test_normalize_date_ignores_the_time_part(raw, iso):
    assert normalize_date(raw) == iso


@pytest.mark.offline
def test_timestamp_periods_and_impossible_dates():
    assert parse_period("2025-01-01 00:00:00") == ((2025, 1), (2025, 1))
    assert parse_period(datetime(2025, 3, 9, 14, 5)) == ((2025, 3), (2025, 3))
    assert parse_period("31/02/2026") is None and parse_period("2026-02-30 00:00:00") is None


@pytest.mark.offline
@pytest.mark.parametrize("timestamp", [True, False])
def test_fuel_records_without_cost_still_give_the_quantity_matrix(timestamp):
    recs = fuel_records(timestamp=timestamp)                    # 6 vehicles x 18 months = 108 records, importo_eur empty
    assert len(recs) == 108 and recs[0]["fields"]["importo_eur"]["status"] == "MISSING"
    a = analyze_records("fleet_fuel_expenses", fuel_schema(), recs)
    assert (a.records_used, a.records_unplaced) == (108, 0)
    assert [(m.field, m.kind, m.unit) for m in a.matrices] == [("litri", "quantity", "L")]   # no cost matrix, no failure
    litri = a.matrices[0]
    assert litri.years == (2025, 2026) and len(litri.by_dimension) == 6
    assert litri.total.values[2025][9] is None                  # months with no records stay empty
    expected = sum(r["fields"]["litri"]["normalized_value"] for r in recs)
    assert litri.grand_total == pytest.approx(expected)


@pytest.mark.offline
def test_records_whose_timestamp_was_stored_as_uncertain_are_still_attributed():
    recs = fuel_records(n_plates=1, years=(2025,), months=range(1, 4))
    for r in recs:                                              # what sessions kept before the normalizer fix
        raw = r["fields"]["data_rifornimento"]["raw_value"]
        r["fields"]["data_rifornimento"] = {"raw_value": raw, "normalized_value": None, "status": "UNCERTAIN"}
    a = analyze_records("fleet_fuel_expenses", fuel_schema(), recs)
    assert a.records_used == 3 and a.matrices[0].years == (2025,)


@pytest.mark.offline
def test_cost_is_added_as_a_second_matrix_when_present():
    a = analyze_records("fleet_fuel_expenses", fuel_schema(), fuel_records(n_plates=2, with_cost=True))
    assert [m.field for m in a.matrices] == ["litri", "importo_eur"]


@pytest.mark.offline
@pytest.mark.parametrize("n_plates,years", [(6, (2025, 2026)), (1, (2025, 2026)), (6, (2025,)), (1, (2025,)),
                                            (len(PALETTE) + 2, (2025, 2026))])
def test_chart_renders_for_every_years_by_dimensions_combination(n_plates, years):
    a = analyze_records("fleet_fuel_expenses", fuel_schema(), fuel_records(n_plates=n_plates, years=years))
    png = render_monthly_chart(a.matrices[0], dpi=60)
    assert png is not None and png.getvalue().startswith(b"\x89PNG")
    assert Image.open(io.BytesIO(png.getvalue())).size == (600, 276)


def _single_module_workbook(records):
    data = {"fuel_expense": {"records": records}}
    out = create_styled_excel_report(data, ["foglio.xlsx"], {"company_name": "X"}, module_id="fleet_fuel_expenses")
    return openpyxl.load_workbook(io.BytesIO(out))


@pytest.mark.offline
def test_single_module_excel_has_the_trend_sheet_with_chart():
    wb = _single_module_workbook(fuel_records(n_plates=3))
    assert TREND_SHEET in wb.sheetnames
    ws = wb[TREND_SHEET]
    assert len(ws._images) == 1                                  # litri only: no cost values, no cost chart
    labels = {c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    assert {"Totale", "Media mensile", "Gennaio", "Targa"} <= labels
    assert wb.sheetnames[0] == "RIEPILOGO_GENERALE"              # existing sheets untouched


@pytest.mark.offline
def test_single_module_excel_without_consumption_data_has_no_trend_sheet():
    recs = fuel_records(n_plates=1, months=range(1, 3))
    for r in recs:
        r["fields"]["litri"] = {"raw_value": None, "normalized_value": None, "status": "MISSING"}
    assert TREND_SHEET not in _single_module_workbook(recs).sheetnames
    assert TREND_SHEET not in _single_module_workbook([]).sheetnames


@pytest.mark.offline
def test_water_bills_route_to_the_utility_schema_and_electricity_stays_on_energy():
    water = [{"text": "Bolletta acquedotto - servizio idrico integrato. Consumo mc 45 fognatura depurazione "
                      "periodo fatturato gennaio 2025 totale da pagare"}]
    h2o = [{"text": "Fattura H2O consumi idrici lettura contatore m3 fornitura acqua potabile"}]
    power = [{"text": "Bolletta energia elettrica POD IT001E123 consumo kWh periodo fatturato totale da pagare"}]
    for pages in (water, h2o):
        for preferred in ("energy_utility", "iso_14001"):
            assert classify_document("doc.pdf", pages, preferred).route == ("iso_14001", "utility_consumption")
    assert classify_document("doc.pdf", power, "iso_14001").route == ("energy_utility", "energy_bill")
