import io
from datetime import date

import openpyxl
import pytest

from app.alerts.dispatcher import dispatch_alerts
from app.alerts.store import (
    add_recipient,
    get_active_deadlines,
    upsert_deadline,
    deactivate_missing_deadlines,
    deactivate_deadlines_for_document,
)
from app.extraction.pipeline import (
    apply_results,
    build_routing_plan,
    classify_document,
    collect_payload,
    evict_document,
    pending_tasks,
    run_batch_extraction,
    sync_deadlines_to_db,
)
from app.reporting.platform_excel_exporter import create_platform_excel_report

BILL = [{"page_number": 1, "text": "Bolletta energia elettrica POD IT001E05324903 consumo 877 kWh. Totale da pagare 250,00"}]
CERT = [{"page_number": 1, "text": "Attestato di formazione - corso antincendio, Mario Rossi, D.Lgs. 81/08"}]
MISC = [{"page_number": 1, "text": "Lettera di cortesia senza particolari riferimenti."}]


@pytest.mark.offline
def test_routing_is_deterministic_and_overridable_by_tiebreak():
    assert classify_document("a.pdf", BILL, "iso_14001").route == ("energy_utility", "energy_bill")
    # exact tie between hr_training and iso_14001 -> active module wins, else registry order
    assert classify_document("c.pdf", CERT, "iso_14001").route == ("iso_14001", "personnel_training")
    assert classify_document("c.pdf", CERT, None).route[1] in ("training_certificate", "personnel_training")
    assert classify_document("c.pdf", CERT, "iso_14001").reason == "tie-resolved"
    # nothing matches -> general_documents fallback
    assert classify_document("x.pdf", MISC).route == ("general_documents", "general")
    assert classify_document("manutenzione_pressa.xlsx", MISC).route == ("machinery_maintenance", "maintenance_log")


def _fake_extract(pages, schema, module_id=None, source_document=""):
    if source_document == "bad.pdf":
        raise RuntimeError("boom")
    return {
        "records": [{
            "source_document": source_document,
            "source_page": 1,
            "warnings": [],
            "fields": {
                "codice_pod": {"raw_value": "IT1", "normalized_value": "IT1", "status": "EXTRACTED"},
                "data_scadenza": {"raw_value": "2026-11-30", "normalized_value": "2026-11-30", "status": "EXTRACTED"},
            },
        }]
    }


@pytest.mark.offline
def test_batch_run_merge_idempotence_and_eviction():
    route = ("energy_utility", "energy_bill")
    extracted, reviewed, versions, routed = {}, {}, {}, {}
    pages = {"a.pdf": BILL, "bad.pdf": BILL}
    tasks = pending_tasks({"a.pdf": route, "bad.pdf": route, "skip.pdf": None}, extracted)
    assert len(tasks) == 2
    seen = []
    res = run_batch_extraction(tasks, pages, on_progress=lambda *a: seen.append(a), extract_fn=_fake_extract)
    apply_results(extracted, reviewed, versions, routed, res)
    assert len(seen) == 2
    assert extracted["energy_utility"]["energy_bill"]["total_records"] == 1
    assert "bad.pdf" in extracted["energy_utility"]["energy_bill"]["errors"]
    # idempotent: successful doc is not re-run, failed doc is retried
    assert pending_tasks({"a.pdf": route, "bad.pdf": route}, extracted) == [("bad.pdf", *route)]
    assert versions[route] == 1 and routed["a.pdf"] == route

    n = sync_deadlines_to_db(collect_payload(extracted, reviewed), ["a.pdf", "bad.pdf"])
    assert n == 1 and len(get_active_deadlines()) == 1
    # removing the document removes its records and stops monitoring
    evict_document(extracted, reviewed, versions, "a.pdf")
    assert extracted["energy_utility"]["energy_bill"]["total_records"] == 0
    assert sync_deadlines_to_db(collect_payload(extracted, reviewed), ["bad.pdf"]) == 0
    deactivate_deadlines_for_document("a.pdf")
    assert get_active_deadlines() == []


@pytest.mark.offline
def test_stale_deadline_deactivated_when_record_changes():
    upsert_deadline("d1", "m", "s", "f", "old", "2026-12-01", [30], "a.pdf", 1)
    upsert_deadline("d2", "m", "s", "f", "other doc", "2026-12-01", [30], "z.pdf", 1)
    assert deactivate_missing_deadlines({"new"}, ["a.pdf"]) == 1
    assert [d["id"] for d in get_active_deadlines()] == ["d2"]


@pytest.mark.offline
def test_unconfigured_notifier_is_not_logged_as_sent(monkeypatch):
    for k in ("SMTP_HOST", "TWILIO_ACCOUNT_SID", "TWILIO_AUTH_TOKEN", "TWILIO_FROM"):
        monkeypatch.delenv(k, raising=False)
    upsert_deadline("d1", "m", "s", "f", "Test", "2026-10-15", [30, 15, 7], "a.pdf", 1)
    add_recipient("r1", "email", "a@b.it")
    stats = dispatch_alerts(today=date(2026, 10, 10), dry_run=False)
    assert stats["simulated"] >= 1 and stats["sent_email"] == 0 and stats["failed"] == 0
    from app.alerts.store import is_alert_already_sent
    assert not is_alert_already_sent("d1", "2026-10-15", 15, "email", "a@b.it")


@pytest.mark.offline
def test_exporter_explains_unanalyzed_documents():
    xl = create_platform_excel_report(
        [{"filename": "empty.pdf"}, {"filename": "bad.pdf"}, {"filename": "new.pdf"}],
        {},
        routing={"empty.pdf": "Energia › Bollette"},
        errors={"bad.pdf": "boom"},
    )
    ws = openpyxl.load_workbook(io.BytesIO(xl))["Registro Documenti"]
    notes = {ws.cell(r, 6).value: ws.cell(r, 13).value for r in range(2, 5)}
    assert "nessun record" in notes["empty.pdf"] and "fallita" in notes["bad.pdf"] and "analisi automatica" in notes["new.pdf"]