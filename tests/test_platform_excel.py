import io
import openpyxl
import pytest
from app.reporting.platform_excel_exporter import create_platform_excel_report


@pytest.mark.offline
def test_platform_excel_4_sheets_structure():
    sample_docs = [{"filename": "bolletta_test.pdf"}]
    sample_extracted_all = {
        "energy_utility": {
            "energy_bill": {
                "records": [{
                    "source_document": "bolletta_test.pdf",
                    "source_page": 1,
                    "fields": {
                        "codice_pod": {"raw_value": "IT001E123", "normalized_value": "IT001E123", "status": "EXTRACTED"},
                        "consumo_kwh": {"raw_value": "1000", "normalized_value": 1000.0, "status": "EXTRACTED"},
                        "data_scadenza": {"raw_value": "2026-11-30", "normalized_value": "2026-11-30", "status": "EXTRACTED"}
                    }
                }]
            }
        }
    }
    company_prof = {"company_name": "Studio & Impresa S.r.l.", "address": "Torino"}

    excel_bytes = create_platform_excel_report(
        documents=sample_docs,
        extracted_data_all_modules=sample_extracted_all,
        company_profile=company_prof
    )

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Dashboard & KPI" in wb.sheetnames
    assert "Registro Documenti" in wb.sheetnames
    assert "Scadenze & Azioni" in wb.sheetnames
    assert "Dati Energy Bill" in wb.sheetnames


@pytest.mark.offline
def test_platform_excel_regulatory_authorization_sheet_title_and_columns():
    sample_docs = [{"filename": "104-2531_2023.pdf"}]
    sample_extracted_all = {
        "iso_14001": {
            "regulatory_authorization": {
                "records": [{
                    "source_document": "104-2531_2023.pdf",
                    "source_page": 1,
                    "fields": {
                        "numero_atto_protocollo": {"raw_value": "104-2531", "normalized_value": "104-2531", "status": "EXTRACTED"},
                        "ente_rilascio": {"raw_value": "Provincia", "normalized_value": "Provincia", "status": "EXTRACTED"},
                        "elenco_codici_eer": {"raw_value": "15 01 01, 17 04 05", "normalized_value": "15 01 01, 17 04 05", "status": "EXTRACTED"},
                        "quantitativi_movimentabili": {"raw_value": "Fase 1: 3000 t/a", "normalized_value": "Fase 1: 3000 t/a", "status": "EXTRACTED"},
                        "quantitativi_deposito": {"raw_value": "Fase 1: 150 t", "normalized_value": "Fase 1: 150 t", "status": "EXTRACTED"},
                        "prescrizioni_con_scadenza": {"raw_value": "Invio piano || entro 90 giorni", "normalized_value": "Invio piano || entro 90 giorni", "status": "EXTRACTED"},
                        "prescrizioni_tecnico_gestionali": {"raw_value": "Stoccaggio su pavimentazione impermeabile", "normalized_value": "Stoccaggio su pavimentazione impermeabile", "status": "EXTRACTED"},
                    }
                }]
            }
        }
    }

    excel_bytes = create_platform_excel_report(
        documents=sample_docs,
        extracted_data_all_modules=sample_extracted_all,
        company_profile={"company_name": "inambiente S.r.l.", "address": "Torino"}
    )

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Dettagli File" in wb.sheetnames
    assert "Dati Regulatory Authorization" not in wb.sheetnames

    ws = wb["Dettagli File"]
    headers = [cell.value for cell in ws[1]]
    idx_eer = headers.index("Codici EER / Rifiuti Ammessi")
    assert headers[idx_eer + 1] == "Quantitativi autorizzati movimentabili"
    assert headers[idx_eer + 2] == "Quantitativi autorizzati in deposito"

    col_psc = headers.index("Prescrizioni con scadenza") + 1
    assert ws.cell(row=2, column=col_psc).alignment.wrap_text is True
    assert ws.cell(row=2, column=col_psc).alignment.vertical == "top"
